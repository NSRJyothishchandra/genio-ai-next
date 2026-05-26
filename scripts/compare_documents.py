import json
import re
import sys
import unicodedata
from dataclasses import dataclass
from html import escape
from itertools import zip_longest
from pathlib import Path
from typing import List

try:
    import fitz as _fitz
    _HAS_FITZ = True
except Exception:
    _fitz = None
    _HAS_FITZ = False

from docx import Document
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

def _highlight_pdf_if_available(left_path, right_path, output_pdf):
    try:
        from pdf_highlight_compare import highlight_pdf
        highlight_pdf(left_path, right_path, output_pdf)
        return True
    except Exception:
        return False


TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".tsv", ".json", ".js", ".jsx", ".ts", ".tsx",
    ".py", ".java", ".cs", ".html", ".htm", ".css", ".scss", ".xml",
    ".yml", ".yaml", ".sql", ".log", ".ini", ".cfg"
}


@dataclass
class Part:
    text: str
    type: str


def read_text_file(file_path: Path) -> str:
    raw = file_path.read_bytes()
    for encoding in ("utf-8", "utf-16", "utf-16-le", "utf-16-be", "latin-1"):
        try:
            return raw.decode(encoding)
        except Exception:
            continue
    raise ValueError(f"Unable to decode text file: {file_path.name}")


def read_docx_file(file_path: Path) -> str:
    document = Document(str(file_path))
    blocks: List[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            blocks.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))

    return "\n".join(blocks)


def _read_pdf_fitz(file_path: Path) -> str:
    doc = _fitz.open(str(file_path))
    pages: List[str] = []
    for page in doc:
        blocks = page.get_text("blocks", sort=True)
        lines: List[str] = []
        for block in blocks:
            if block[6] != 0:  # skip image blocks
                continue
            text = block[4].strip()
            if text:
                lines.append(text)
        page_text = "\n".join(lines)
        if page_text:
            pages.append(page_text)
    doc.close()
    return "\n\n".join(pages)


def _read_pdf_pypdf(file_path: Path) -> str:
    from pypdf import PdfReader
    reader = PdfReader(str(file_path))
    pages: List[str] = []
    for page in reader.pages:
        text = (page.extract_text() or "").strip()
        if text:
            pages.append(text)
    return "\n\n".join(pages)


def read_pdf_file(file_path: Path) -> str:
    if _HAS_FITZ:
        return _read_pdf_fitz(file_path)
    return _read_pdf_pypdf(file_path)


def _normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFC", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text


def extract_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return _normalize_text(read_pdf_file(file_path))
    if suffix == ".docx":
        return _normalize_text(read_docx_file(file_path))
    if suffix in TEXT_EXTENSIONS:
        text = read_text_file(file_path)
        if suffix == ".json":
            try:
                parsed = json.loads(text)
                return _normalize_text(json.dumps(parsed, indent=2, ensure_ascii=False))
            except Exception:
                return _normalize_text(text)
        return _normalize_text(text)
    raise ValueError(
        f"Unsupported file type: {file_path.suffix or 'unknown'}. "
        "Supported types include DOCX, TXT, MD, CSV, JSON, code, XML, YAML, and log files."
    )


def tokenize(text: str) -> List[str]:
    return re.findall(r"\s+|\w+|[^\w\s]", text, flags=re.UNICODE)


def merge_parts(parts: List[Part]) -> List[dict]:
    merged: List[Part] = []
    for part in parts:
        if not part.text:
            continue
        if merged and merged[-1].type == part.type:
            merged[-1].text += part.text
        else:
            merged.append(Part(part.text, part.type))
    return [{"text": part.text, "type": part.type} for part in merged]


def _lines_are_similar(a: str, b: str, threshold: float = 0.30) -> bool:
    import difflib
    if not a and not b:
        return True
    if not a or not b:
        return False
    return difflib.SequenceMatcher(None, a, b, autojunk=False).ratio() >= threshold


def build_intraline_parts(left: str, right: str):
    import difflib

    left_parts: List[Part] = []
    right_parts: List[Part] = []
    matcher = difflib.SequenceMatcher(None, tokenize(left), tokenize(right), autojunk=False)

    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        left_text = "".join(tokenize(left)[i1:i2])
        right_text = "".join(tokenize(right)[j1:j2])
        if tag == "equal":
            left_parts.append(Part(left_text, "equal"))
            right_parts.append(Part(right_text, "equal"))
        elif tag == "delete":
            left_parts.append(Part(left_text, "delete"))
        elif tag == "insert":
            right_parts.append(Part(right_text, "insert"))
        else:
            left_parts.append(Part(left_text, "delete"))
            right_parts.append(Part(right_text, "insert"))

    return merge_parts(left_parts), merge_parts(right_parts)


def make_plain_parts(text: str) -> List[dict]:
    return [{"text": text, "type": "equal"}]


def build_entries(left_text: str, right_text: str) -> List[dict]:
    import difflib

    left_lines = left_text.splitlines()
    right_lines = right_text.splitlines()
    matcher = difflib.SequenceMatcher(None, left_lines, right_lines, autojunk=False)
    entries: List[dict] = []
    left_number = 1
    right_number = 1

    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        if tag == "equal":
            for left_line, right_line in zip(left_lines[i1:i2], right_lines[j1:j2]):
                entries.append({
                    "kind": "equal",
                    "leftLineNumber": left_number,
                    "rightLineNumber": right_number,
                    "leftParts": make_plain_parts(left_line),
                    "rightParts": make_plain_parts(right_line),
                })
                left_number += 1
                right_number += 1
            continue

        if tag == "replace":
            left_block = left_lines[i1:i2]
            right_block = right_lines[j1:j2]
            for left_line, right_line in zip_longest(left_block, right_block, fillvalue=""):
                if left_line and right_line:
                    # Always keep both sides on the same row so the table shows
                    # original vs. updated values for every changed line.
                    if _lines_are_similar(left_line, right_line):
                        left_parts, right_parts = build_intraline_parts(left_line, right_line)
                    else:
                        left_parts = [{"text": left_line, "type": "delete"}]
                        right_parts = [{"text": right_line, "type": "insert"}]
                    entries.append({
                        "kind": "replace",
                        "leftLineNumber": left_number,
                        "rightLineNumber": right_number,
                        "leftParts": left_parts,
                        "rightParts": right_parts,
                    })
                    left_number += 1
                    right_number += 1
                elif left_line:
                    entries.append({
                        "kind": "delete",
                        "leftLineNumber": left_number,
                        "rightLineNumber": None,
                        "leftParts": [{"text": left_line, "type": "delete"}],
                        "rightParts": [],
                    })
                    left_number += 1
                elif right_line:
                    entries.append({
                        "kind": "insert",
                        "leftLineNumber": None,
                        "rightLineNumber": right_number,
                        "leftParts": [],
                        "rightParts": [{"text": right_line, "type": "insert"}],
                    })
                    right_number += 1
            continue

        if tag == "delete":
            for left_line in left_lines[i1:i2]:
                entries.append({
                    "kind": "delete",
                    "leftLineNumber": left_number,
                    "rightLineNumber": None,
                    "leftParts": [{"text": left_line, "type": "delete"}],
                    "rightParts": [],
                })
                left_number += 1
            continue

        if tag == "insert":
            for right_line in right_lines[j1:j2]:
                entries.append({
                    "kind": "insert",
                    "leftLineNumber": None,
                    "rightLineNumber": right_number,
                    "leftParts": [],
                    "rightParts": [{"text": right_line, "type": "insert"}],
                })
                right_number += 1

    return entries


def build_summary(entries: List[dict], left_text: str, right_text: str) -> dict:
    counts = {"equal": 0, "replace": 0, "insert": 0, "delete": 0}
    for entry in entries:
        counts[entry["kind"]] += 1
    return {
        "leftLineCount": len(left_text.splitlines()),
        "rightLineCount": len(right_text.splitlines()),
        "equalLines": counts["equal"],
        "changedLines": counts["replace"] + counts["insert"] + counts["delete"],
        "replacedLines": counts["replace"],
        "insertedLines": counts["insert"],
        "deletedLines": counts["delete"],
    }


def parts_to_markup(parts: List[dict]) -> str:
    if not parts:
        return ""

    markup = []
    for part in parts:
        text = (
            escape(part["text"])
            .replace("\t", "&nbsp;&nbsp;&nbsp;&nbsp;")
            .replace(" ", "&nbsp;")
        )
        if part["type"] == "insert":
            markup.append(f'<font backColor="#D1FAE5">{text}</font>')
        elif part["type"] == "delete":
            markup.append(f'<font backColor="#FECACA">{text}</font>')
        else:
            markup.append(text)
    return "".join(markup) or "&nbsp;"


def build_context_entries(entries: List[dict], context_radius: int = 2) -> List[dict]:
    changed_indexes = [index for index, entry in enumerate(entries) if entry["kind"] != "equal"]
    if not changed_indexes:
        return entries[: min(len(entries), 100)]

    selected = set()
    for index in changed_indexes:
        for offset in range(-context_radius, context_radius + 1):
            candidate = index + offset
            if 0 <= candidate < len(entries):
                selected.add(candidate)

    context_entries = []
    previous = None
    for index in sorted(selected):
        if previous is not None and index - previous > 1:
            context_entries.append({"kind": "gap"})
        context_entries.append(entries[index])
        previous = index
    return context_entries


def build_updated_focus_entries(entries: List[dict]) -> List[dict]:
    focused = []
    for entry in entries:
        if entry["kind"] == "equal":
            continue
        focused.append({
            "kind": entry["kind"],
            "lineNumber": entry["rightLineNumber"],
            "parts": entry["rightParts"] if entry["rightParts"] else [{"text": "", "type": "equal"}],
        })
    return focused


def build_pdf_report(output_pdf: Path, report: dict):
    output_pdf.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(
        str(output_pdf),
        pagesize=landscape(A4),
        leftMargin=10 * mm,
        rightMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )

    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    body_style = ParagraphStyle(
        "Body",
        parent=styles["BodyText"],
        fontName="Courier",
        fontSize=7.4,
        leading=9,
        allowWidows=1,
        allowOrphans=1,
    )
    small_style = ParagraphStyle("Small", parent=styles["BodyText"], fontSize=8, leading=10)
    line_no_style = ParagraphStyle("LineNo", parent=small_style, fontName="Courier-Bold", alignment=1)

    story = [
        Paragraph("Document Difference Report", title_style),
        Spacer(1, 6),
        Paragraph(
            f"<b>Original file:</b> {escape(report['leftFileName'])} &nbsp;&nbsp;&nbsp; "
            f"<b>Updated file:</b> {escape(report['rightFileName'])}",
            small_style,
        ),
        Spacer(1, 6),
    ]

    summary_data = [
        ["Original lines", str(report["summary"]["leftLineCount"]), "Updated lines", str(report["summary"]["rightLineCount"])],
        ["Changed lines", str(report["summary"]["changedLines"]), "Equal lines", str(report["summary"]["equalLines"])],
        ["Inserted", str(report["summary"]["insertedLines"]), "Deleted", str(report["summary"]["deletedLines"])],
        ["Replaced", str(report["summary"]["replacedLines"]), "Generated at", escape(report["generatedAt"])],
    ]
    summary_table = Table(summary_data, colWidths=[30 * mm, 35 * mm, 30 * mm, 95 * mm])
    summary_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.whitesmoke),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    story.extend([summary_table, Spacer(1, 10)])

    story.append(Paragraph("Detailed highlighted comparison", styles["Heading2"]))
    story.append(Spacer(1, 4))

    rows = [[
        Paragraph("<b>L#</b>", small_style),
        Paragraph("<b>Original</b>", small_style),
        Paragraph("<b>R#</b>", small_style),
        Paragraph("<b>Updated</b>", small_style),
    ]]
    table_styles = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E2E8F0")),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]

    context_entries = build_context_entries(report["entries"], context_radius=2)
    for row_index, entry in enumerate(context_entries, start=1):
        if entry.get("kind") == "gap":
            rows.append([
                Paragraph("", line_no_style),
                Paragraph("<i>... unchanged section omitted ...</i>", small_style),
                Paragraph("", line_no_style),
                Paragraph("<i>... unchanged section omitted ...</i>", small_style),
            ])
            table_styles.append(("BACKGROUND", (0, row_index), (-1, row_index), colors.HexColor("#F8FAFC")))
            continue

        left_number = "" if entry["leftLineNumber"] is None else str(entry["leftLineNumber"])
        right_number = "" if entry["rightLineNumber"] is None else str(entry["rightLineNumber"])
        left_markup = parts_to_markup(entry["leftParts"])
        right_markup = parts_to_markup(entry["rightParts"])
        rows.append([
            Paragraph(left_number or "&nbsp;", line_no_style),
            Paragraph(left_markup, body_style),
            Paragraph(right_number or "&nbsp;", line_no_style),
            Paragraph(right_markup, body_style),
        ])

        if entry["kind"] == "insert":
            table_styles.append(("BACKGROUND", (2, row_index), (3, row_index), colors.HexColor("#ECFDF5")))
        elif entry["kind"] == "delete":
            table_styles.append(("BACKGROUND", (0, row_index), (1, row_index), colors.HexColor("#FEF2F2")))
        elif entry["kind"] == "replace":
            table_styles.append(("BACKGROUND", (0, row_index), (1, row_index), colors.HexColor("#FEF2F2")))
            table_styles.append(("BACKGROUND", (2, row_index), (3, row_index), colors.HexColor("#ECFDF5")))
        else:
            table_styles.append(("BACKGROUND", (0, row_index), (-1, row_index), colors.white))

    comparison_table = Table(rows, colWidths=[12 * mm, 120 * mm, 12 * mm, 120 * mm], repeatRows=1)
    comparison_table.setStyle(TableStyle(table_styles))
    story.append(comparison_table)
    story.append(PageBreak())

    story.append(Paragraph("Updated File Highlighted Changes", styles["Heading2"]))
    story.append(Spacer(1, 4))
    updated_rows = [[
        Paragraph("<b>Updated line</b>", small_style),
        Paragraph("<b>Updated file content with highlights</b>", small_style),
    ]]
    updated_styles = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DBEAFE")),
        ("BOX", (0, 0), (-1, -1), 0.5, colors.lightgrey),
        ("INNERGRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]

    for row_index, entry in enumerate(build_updated_focus_entries(report["entries"]), start=1):
        updated_rows.append([
            Paragraph("" if entry["lineNumber"] is None else str(entry["lineNumber"]), line_no_style),
            Paragraph(parts_to_markup(entry["parts"]), body_style),
        ])
        updated_styles.append(("BACKGROUND", (0, row_index), (-1, row_index), colors.HexColor("#ECFDF5")))

    if len(updated_rows) == 1:
        updated_rows.append([
            Paragraph("", line_no_style),
            Paragraph("No updated differences were found.", small_style),
        ])

    updated_table = Table(updated_rows, colWidths=[20 * mm, 244 * mm], repeatRows=1)
    updated_table.setStyle(TableStyle(updated_styles))
    story.append(updated_table)
    story.append(Spacer(1, 10))

    story.append(Paragraph("Comparison notes", styles["Heading2"]))
    story.append(Spacer(1, 4))
    story.append(Paragraph(
        "Green highlights show inserted or changed content in the updated file. Red highlights show removed or replaced content from the original file. "
        "This report is optimized for textual precision across PDF, DOCX, and text-based files.",
        small_style,
    ))

    doc.build(story)


def main():
    if len(sys.argv) != 5:
        print("Usage: compare_documents.py <left_file> <right_file> <output_json> <output_pdf>", file=sys.stderr)
        sys.exit(1)

    left_path = Path(sys.argv[1])
    right_path = Path(sys.argv[2])
    output_json = Path(sys.argv[3])
    output_pdf = Path(sys.argv[4])
    both_pdf = left_path.suffix.lower() == ".pdf" and right_path.suffix.lower() == ".pdf"

    left_text = extract_text(left_path).replace("\r\n", "\n").replace("\r", "\n")
    right_text = extract_text(right_path).replace("\r\n", "\n").replace("\r", "\n")

    entries = build_entries(left_text, right_text)
    report = {
        "leftFileName": left_path.name,
        "rightFileName": right_path.name,
        "generatedAt": __import__("datetime").datetime.now().isoformat(),
        "summary": build_summary(entries, left_text, right_text),
        "entries": entries,
    }

    output_json.parent.mkdir(parents=True, exist_ok=True)
    output_json.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    if both_pdf and _highlight_pdf_if_available(left_path, right_path, output_pdf):
        pass
    else:
        build_pdf_report(output_pdf, report)
    print(json.dumps({
        "jsonPath": str(output_json),
        "pdfPath": str(output_pdf),
        "summary": report["summary"],
    }))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        sys.exit(1)
